"""Tests for KPIs, breakdowns, growth, customers, forecast, insights and export."""
import io
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src import analytics, exporter
from src.data_cleaning import clean_data
from src.generate_sample_data import generate_sales_data
from tests.helpers import clean_frame, two_year_rows


def synthetic_clean(seed=3):
    raw = generate_sales_data(n_orders=2500, n_customers=200, seed=seed, add_dirty_rows=False)
    return clean_data(raw.astype(str))[0]


class KpiTests(unittest.TestCase):
    def test_kpis_match_hand_calculation(self):
        k = analytics.calculate_kpis(clean_frame())
        self.assertAlmostEqual(k["total_revenue"], 595.0)
        self.assertEqual(k["total_orders"], 3)               # distinct orders, not 4 rows
        self.assertEqual(k["total_units"], 7)
        self.assertAlmostEqual(k["average_order_value"], 595 / 3)
        self.assertEqual(k["unique_customers"], 2)
        self.assertAlmostEqual(k["gross_profit"], 215.0)
        self.assertAlmostEqual(k["profit_margin_pct"], 215 / 595 * 100)

    def test_empty_dataframe_gives_zero_kpis(self):
        k = analytics.calculate_kpis(clean_frame().iloc[0:0])
        self.assertEqual(k["total_revenue"], 0.0)
        self.assertIsNone(k["gross_profit"])

    def test_kpis_agree_with_independent_numpy_calculation(self):
        df = synthetic_clean()
        manual = float(np.sum(df["quantity"].to_numpy() * df["unit_price"].to_numpy()
                              * (1 - df["discount"].to_numpy())))
        self.assertAlmostEqual(analytics.calculate_kpis(df)["total_revenue"], manual, places=2)


class BreakdownTests(unittest.TestCase):
    def test_every_dimension_sums_to_total_revenue(self):
        df = synthetic_clean()
        total = df["revenue"].sum()
        for dim in ("category", "product", "region"):
            table = analytics.sales_by_dimension(df, dim)
            self.assertAlmostEqual(table["revenue"].sum(), total, places=2)
            self.assertAlmostEqual(table["revenue_share_pct"].sum(), 100.0, places=6)

    def test_pareto_ends_at_100_and_marks_core_items(self):
        pareto = analytics.pareto_analysis(synthetic_clean(), "product")
        self.assertAlmostEqual(pareto["cumulative_share_pct"].iloc[-1], 100.0, places=6)
        self.assertTrue(pareto["is_core_contributor"].iloc[0])
        self.assertFalse(pareto["is_core_contributor"].iloc[-1])

    def test_top_and_bottom(self):
        top, bottom = analytics.top_and_bottom(synthetic_clean(), "product", 3)
        self.assertEqual(len(top), 3)
        self.assertGreaterEqual(top["revenue"].iloc[0], top["revenue"].iloc[-1])
        self.assertLessEqual(bottom["revenue"].iloc[0], bottom["revenue"].iloc[-1])

    def test_dimension_table_on_empty_data(self):
        out = analytics.sales_by_dimension(clean_frame().iloc[0:0], "category")
        self.assertTrue(out.empty)


class FilterTests(unittest.TestCase):
    def test_filters_combine_and_end_date_is_inclusive(self):
        df = clean_frame()
        out = analytics.apply_filters(df, "2024-01-10", "2024-02-15", regions=["South"])
        self.assertEqual(out["order_id"].tolist(), ["O2"])
        both = analytics.apply_filters(df, "2024-01-10", "2024-02-15")
        self.assertEqual(set(both["order_id"]), {"O1", "O2"})

    def test_none_means_no_filter_and_empty_list_means_no_rows(self):
        df = clean_frame()
        self.assertEqual(len(analytics.apply_filters(df)), len(df))
        self.assertTrue(analytics.apply_filters(df, regions=[]).empty)


class TimeSeriesTests(unittest.TestCase):
    def test_month_over_month_growth(self):
        m = analytics.monthly_sales(clean_frame())
        self.assertEqual(m["revenue"].round(2).tolist(), [245.0, 200.0, 150.0])
        self.assertAlmostEqual(m["mom_growth_pct"].iloc[1], (200 - 245) / 245 * 100)
        self.assertTrue(np.isnan(m["mom_growth_pct"].iloc[0]))

    def test_year_over_year_growth_uses_same_month_last_year(self):
        m = analytics.monthly_sales(clean_frame(two_year_rows()))
        self.assertEqual(len(m), 14)
        self.assertAlmostEqual(m["yoy_growth_pct"].iloc[12], 50.0)   # Jan 2024 vs Jan 2023
        self.assertAlmostEqual(m["yoy_growth_pct"].iloc[13], 50.0)

    def test_missing_months_are_filled_with_zero(self):
        rows = [r for r in two_year_rows() if not r[1].startswith("2023-05")]
        m = analytics.monthly_sales(clean_frame(rows))
        self.assertEqual(len(m), 14)
        self.assertEqual(m.loc[m["month"] == "2023-05-01", "revenue"].iloc[0], 0.0)

    def test_incomplete_last_month_is_dropped(self):
        m = analytics.monthly_sales(clean_frame())
        self.assertEqual(len(analytics.drop_incomplete_last_month(m, "2024-03-20")), 2)
        self.assertEqual(len(analytics.drop_incomplete_last_month(m, "2024-03-31")), 3)

    def test_period_tables(self):
        q = analytics.sales_by_period(clean_frame(), "quarter")
        self.assertEqual(q["revenue"].round(2).tolist(), [595.0])
        with self.assertRaises(ValueError):
            analytics.sales_by_period(clean_frame(), "decade")

    def test_period_growth_needs_two_windows(self):
        self.assertTrue(analytics.period_growth_by_dimension(clean_frame(), "product").empty)
        full = analytics.period_growth_by_dimension(synthetic_clean(), "product")
        self.assertFalse(full.empty)


class CustomerTests(unittest.TestCase):
    def test_repeat_metrics(self):
        r = analytics.repeat_purchase_metrics(clean_frame())
        self.assertEqual(r["repeat_customers"], 1)            # only C1 has 2 orders
        self.assertAlmostEqual(r["repeat_rate_pct"], 50.0)
        self.assertAlmostEqual(r["repeat_revenue_share_pct"], (200 + 45 + 150) / 595 * 100)

    def test_rfm_covers_every_customer_with_valid_scores(self):
        df = synthetic_clean()
        rfm = analytics.rfm_segmentation(df)
        self.assertEqual(len(rfm), df["customer_id"].nunique())
        self.assertTrue(rfm["segment"].notna().all())
        for col in ("r_score", "f_score", "m_score"):
            self.assertTrue(rfm[col].between(1, 5).all())
        summary = analytics.rfm_segment_summary(rfm)
        self.assertAlmostEqual(summary["revenue_share_pct"].sum(), 100.0, places=6)

    def test_rfm_on_empty_data(self):
        self.assertTrue(analytics.rfm_segmentation(clean_frame().iloc[0:0]).empty)


class ForecastTests(unittest.TestCase):
    def test_forecast_shape_and_non_negative(self):
        m = analytics.monthly_sales(synthetic_clean())
        res = analytics.forecast_revenue(m, horizon=6)
        self.assertTrue(res["available"])
        self.assertEqual(len(res["forecast"]), 6)
        self.assertTrue((res["forecast"]["forecast"] >= 0).all())
        self.assertTrue((res["forecast"]["lower"] <= res["forecast"]["upper"]).all())
        self.assertEqual(res["forecast"]["month"].iloc[0], m["month"].iloc[-1] + pd.offsets.MonthBegin(1))
        self.assertIsNotNone(res["metrics"])

    def test_refuses_to_forecast_with_too_little_data(self):
        m = analytics.monthly_sales(clean_frame())
        res = analytics.forecast_revenue(m)
        self.assertFalse(res["available"])
        self.assertIn("at least", res["message"])

    def test_perfect_linear_series_is_forecast_accurately(self):
        months = pd.date_range("2022-01-01", periods=30, freq="MS")
        m = pd.DataFrame({"month": months, "revenue": 1000.0 + 50.0 * np.arange(30)})
        res = analytics.forecast_revenue(m, horizon=3)
        self.assertAlmostEqual(res["forecast"]["forecast"].iloc[0], 1000 + 50 * 30, places=4)
        self.assertLess(res["metrics"]["model_mape"], 0.01)


class InsightAndExportTests(unittest.TestCase):
    def test_insights_are_non_empty_and_handle_empty_data(self):
        out = analytics.generate_insights(synthetic_clean())
        self.assertGreater(len(out["findings"]), 3)
        self.assertTrue(all(isinstance(s, str) and s for s in out["findings"]))
        empty = analytics.generate_insights(clean_frame().iloc[0:0])
        self.assertEqual(empty["recommendations"], [])

    def test_excel_report_has_expected_sheets_and_synthetic_label(self):
        data = exporter.build_excel_report(clean_frame(), "test filters", is_synthetic=True)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Summary", "By Category", "By Product",
                                         "By Region", "Monthly", "Filtered Data"])
        summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
        self.assertTrue(any("SYNTHETIC" in str(r[1]) for r in summary))
        self.assertEqual(wb["Filtered Data"].max_row, 1 + 4)

    def test_csv_export_respects_given_rows(self):
        df = clean_frame()
        csv = exporter.dataframe_to_csv_bytes(df[df["region"] == "South"]).decode("utf-8-sig")
        self.assertEqual(len(csv.strip().splitlines()), 2)   # header + 1 row


if __name__ == "__main__":
    unittest.main()
